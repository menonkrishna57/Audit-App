"""
generate_deck.py
Generates an executive 12-slide widescreen PowerPoint presentation for Audit-App
using python-pptx with a modern dark tech aesthetic.
"""

import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path: str = "Audit_App_Presentation.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]  # completely blank slide

    # Theme Palette (Deep Slate & Vibrant Cyber Accents)
    COLOR_BG = RGBColor(15, 23, 42)          # Dark Slate (#0F172A)
    COLOR_CARD = RGBColor(30, 41, 59)        # Card Slate (#1E293B)
    COLOR_CARD_BORDER = RGBColor(51, 65, 85) # Border Slate (#334155)
    COLOR_ACCENT = RGBColor(56, 189, 248)    # Cyan / Sky (#38BDF8)
    COLOR_GREEN = RGBColor(52, 211, 153)     # Emerald (#34D399)
    COLOR_AMBER = RGBColor(251, 191, 36)     # Amber (#FBBF24)
    COLOR_ROSE = RGBColor(251, 113, 133)     # Coral Rose (#FB7185)
    COLOR_WHITE = RGBColor(248, 250, 252)    # Crisp White (#F8FAFC)
    COLOR_MUTED = RGBColor(148, 163, 184)    # Muted Gray (#94A3B8)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="AUDIT-APP PLATFORM OVERVIEW", slide_num=""):
        # Category Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(10), Inches(0.4))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = category_text.upper()
        p_tag.font.name = "Segoe UI"
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_ACCENT

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(10.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Segoe UI"
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_WHITE

        if slide_num:
            num_box = slide.shapes.add_textbox(Inches(11.5), Inches(0.45), Inches(1.2), Inches(0.4))
            tf_num = num_box.text_frame
            p_num = tf_num.paragraphs[0]
            p_num.text = str(slide_num)
            p_num.alignment = PP_ALIGN.RIGHT
            p_num.font.name = "Segoe UI"
            p_num.font.size = Pt(12)
            p_num.font.bold = True
            p_num.font.color.rgb = COLOR_MUTED

    def add_card(slide, left, top, width, height, title, items, badge="", badge_color=COLOR_ACCENT):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_CARD_BORDER
        card.line.width = Pt(1)

        tx_box = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(width - 0.5), Inches(height - 0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True

        p_head = tf.paragraphs[0]
        p_head.text = title
        p_head.font.name = "Segoe UI"
        p_head.font.size = Pt(16)
        p_head.font.bold = True
        p_head.font.color.rgb = COLOR_WHITE
        p_head.space_after = Pt(10)

        for it in items:
            p_item = tf.add_paragraph()
            p_item.font.name = "Segoe UI"
            p_item.font.size = Pt(12)
            p_item.space_after = Pt(6)
            if isinstance(it, tuple):
                bold_prefix, normal_text = it
                run1 = p_item.add_run()
                run1.text = f"• {bold_prefix}: "
                run1.font.bold = True
                run1.font.color.rgb = COLOR_ACCENT
                run2 = p_item.add_run()
                run2.text = normal_text
                run2.font.color.rgb = COLOR_WHITE
            else:
                run = p_item.add_run()
                run.text = f"• {it}"
                run.font.color.rgb = COLOR_WHITE

        return card

    # =========================================================================
    # SLIDE 1: Title Slide (Hero)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Accent decorative glow rectangle
    glow = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.15), Inches(3.6))
    glow.fill.solid()
    glow.fill.fore_color.rgb = COLOR_ACCENT
    glow.line.fill.background()

    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(1.6), Inches(11), Inches(3.8))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "ENTERPRISE SECURITY INTELLIGENCE"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT
    p0.space_after = Pt(12)

    p1 = tf1.add_paragraph()
    p1.text = "Audit-App"
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.space_after = Pt(8)

    p2 = tf1.add_paragraph()
    p2.text = "Next-Generation Autonomous Code Security & Cryptographic Compliance Pipeline"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(20)
    p2.font.color.rgb = COLOR_MUTED
    p2.space_after = Pt(24)

    p3 = tf1.add_paragraph()
    p3.text = "SBOM  |  CBOM & Quantum Posture  |  RCA with Code Fixes  |  OWASP & CWE Gap Analysis"
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(14)
    p3.font.bold = True
    p3.font.color.rgb = COLOR_GREEN

    # Bottom Metadata Pill Card
    b_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(5.6), Inches(10.8), Inches(1.0))
    b_card.fill.solid()
    b_card.fill.fore_color.rgb = COLOR_CARD
    b_card.line.color.rgb = COLOR_CARD_BORDER
    bt_box = s1.shapes.add_textbox(Inches(1.4), Inches(5.7), Inches(10.4), Inches(0.8))
    btf = bt_box.text_frame
    bp = btf.paragraphs[0]
    bp.text = "Architecture: AI-Powered RAG (Google Gemini + Qdrant Vector Engine) + Deterministic CI/CD Static Fallback"
    bp.font.name = "Segoe UI"
    bp.font.size = Pt(12)
    bp.font.color.rgb = COLOR_WHITE
    bp2 = btf.add_paragraph()
    bp2.text = "Target Languages: Python (.py, .ipynb) | Java (.java) | C/C++ (.cpp, .c, .h) | Manifest Ecosystems"
    bp2.font.name = "Segoe UI"
    bp2.font.size = Pt(11)
    bp2.font.color.rgb = COLOR_MUTED

    # =========================================================================
    # SLIDE 2: Problem Statement & Industry Challenges
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "The Code Security Crisis: Why Traditional Tools Fall Short", "INDUSTRY CHALLENGES", "02/12")

    add_card(s2, 0.8, 1.8, 3.6, 5.0, "Supply Chain Blindspots", [
        ("Explosion of Dependencies", "Modern microservices rely on hundreds of third-party libraries; manual tracking is impossible."),
        ("Dependency Confusion", "Unpinned packages allow untrusted upstream releases to introduce zero-days silently."),
        ("Missing SBOMs", "Lack of automated machine-readable bills of materials fails modern compliance audits.")
    ], badge_color=COLOR_ROSE)

    add_card(s2, 4.8, 1.8, 3.6, 5.0, "Cryptographic Obsolescence", [
        ("Quantum Horizon Risk", "Legacy RSA and ECC algorithms will soon be broken by Shor's algorithm; teams lack crypto visibility."),
        ("Hardcoded Ciphers", "MD5 and SHA-1 still linger in repositories for checksums, triggering compliance violations."),
        ("No Native CBOM", "Virtually zero commercial linters build an actionable Cryptography Bill of Materials.")
    ], badge_color=COLOR_AMBER)

    add_card(s2, 8.8, 1.8, 3.6, 5.0, "The 'Scan & Dump' Failure", [
        ("Alert Fatigue", "Traditional SAST scanners dump hundreds of warning alerts without context or triage."),
        ("No Actionable Fixes", "Developers waste hours researching how to fix warnings; linters don't suggest safe replacement code."),
        ("Shallow Regex", "Linters fail to understand multi-file architectural flaws, prompt injection, or root causes.")
    ], badge_color=COLOR_ACCENT)

    # =========================================================================
    # SLIDE 3: The Audit-App Solution (4 Pillars)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Unified 4-Pillar Security & Compliance Framework", "THE SOLUTION", "03/12")

    add_card(s3, 0.8, 1.8, 2.7, 5.0, "1. SBOM Engine", [
        ("Comprehensive Inventory", "Captures third-party libraries across Python, Java, C++, and manifests."),
        ("Version Pinning", "Identifies exact pinned vs. loose range-specified packages."),
        ("Ecosystem Mapping", "Tracks PyPI, Maven, Gradle, and source-level imports.")
    ], badge_color=COLOR_GREEN)

    add_card(s3, 3.8, 1.8, 2.7, 5.0, "2. CBOM Engine", [
        ("Crypto Primitive Catalog", "Discovers all hashes, symmetric ciphers, asymmetric keys, and RNGs."),
        ("Quantum Readiness", "Evaluates quantum vulnerability (RSA/ECC vs. ML-KEM/ML-DSA)."),
        ("Weakness Alerts", "Flags MD5, SHA-1, DES, verify=False, and predictable PRNGs.")
    ], badge_color=COLOR_ACCENT)

    add_card(s3, 6.8, 1.8, 2.7, 5.0, "3. RCA & Code Fixes", [
        ("Root Cause Diagnosis", "Traces why vulnerabilities exist rather than just reporting symptoms."),
        ("Line Number Citing", "Points directly to exact file names and line numbers."),
        ("Drop-in Fix Snippets", "Provides complete recommended replacement code without modifying files.")
    ], badge_color=COLOR_AMBER)

    add_card(s3, 9.8, 1.8, 2.7, 5.0, "4. CWE & OWASP Gaps", [
        ("OWASP Top 10 (2021)", "Evaluates 10 standard categories: Covered, Partial, or Gap."),
        ("CWE Top 25 (2024)", "Maps weaknesses to CWE-798, CWE-78, CWE-89, CWE-502, CWE-22, etc."),
        ("Remediation Roadmap", "Action list prioritized into 4 execution tiers for sprint planning.")
    ], badge_color=COLOR_ROSE)

    # =========================================================================
    # SLIDE 4: Architecture & Workflow
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Under the Hood: Dual-Engine AI RAG & Deterministic Core", "SYSTEM ARCHITECTURE", "04/12")

    add_card(s4, 0.8, 1.8, 5.6, 5.0, "AI RAG Pipeline (Gemini + Qdrant)", [
        ("Line-Preserving Chunker", "Splits source code while maintaining 1-based start_line and end_line chunk metadata."),
        ("768-dim Vector Embeddings", "Generates dense semantic vectors via gemini-embedding-001 or Jina AI code embeddings."),
        ("Qdrant Vector Database", "Stores and retrieves semantic code chunks using Cosine similarity search."),
        ("Multi-Query RAG Search", "Runs targeted domain queries tailored specifically for SBOM, CBOM, RCA, and CWE."),
        ("Gemini Reasoning Engine", "Evaluates retrieved contexts with gemini-3.5-flash to produce structured executive reports.")
    ], badge_color=COLOR_ACCENT)

    add_card(s4, 6.8, 1.8, 5.6, 5.0, "Deterministic Static Fallback Engine", [
        ("Zero-Dependency Execution", "Runs directly on Python built-ins (ast, re, json, pathlib) without external network calls."),
        ("CI/CD & Offline Reliability", "Guarantees full security audit continuity even when API keys or databases are unavailable."),
        ("AST & Pattern Matchers", "Inspects AST trees and language patterns across Python, Java, C/C++, and Jupyter Notebooks."),
        ("Automated Report Assembly", "Compiles full Markdown and JSON artifacts identically structured to the AI reports."),
        ("GitHub Summary Publisher", "Pushes structured tables straight to $GITHUB_STEP_SUMMARY for instant developer review.")
    ], badge_color=COLOR_GREEN)

    # =========================================================================
    # SLIDE 5: Pillar 1 — SBOM (Software Bill of Materials)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Pillar 1: Software Bill of Materials (SBOM)", "DEPENDENCY INTELLIGENCE", "05/12")

    add_card(s5, 0.8, 1.8, 5.6, 5.0, "Multi-Ecosystem Discovery", [
        ("Python Ecosystem", "Parses requirements.txt, pyproject.toml, and unpinned direct source imports."),
        ("Java Ecosystem", "Extracts Maven pom.xml (groupId, artifactId, version) and Gradle build.gradle dependencies."),
        ("C/C++ Ecosystem", "Analyzes #include header directives and CMake dependencies."),
        ("Notebook Extraction", "Specialized Jupyter (.ipynb) parser extracts cell-level imports while stripping base64 image bloat.")
    ], badge_color=COLOR_ACCENT)

    add_card(s5, 6.8, 1.8, 5.6, 5.0, "Actionable Compliance Output", [
        ("Component Inventory Table", "Component name, exact declared version, ecosystem, declaration source, usage context, and risk notes."),
        ("Pinned vs. Unpinned Metrics", "Instantly flags loose version ranges vulnerable to supply-chain tampering."),
        ("Supply-Chain Protection", "Guards against dependency confusion attacks and out-of-date packages."),
        ("Machine-Readable Formats", "Emits standard Markdown tables and JSON payloads for integration into procurement and audit logs.")
    ], badge_color=COLOR_GREEN)

    # =========================================================================
    # SLIDE 6: Pillar 2 — CBOM & Quantum Posture
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Pillar 2: Cryptography Bill of Materials (CBOM)", "QUANTUM READINESS & COMPLIANCE", "06/12")

    add_card(s6, 0.8, 1.8, 5.6, 5.0, "Cryptographic Asset Inventory", [
        ("Primitives Cataloged", "Hashing (MD5, SHA-1, SHA-256), Ciphers (AES, RSA, DES), Protocols (TLS/SSL), RNGs (secrets, random)."),
        ("Key Sizes & Libraries", "Audits exact key bit lengths (e.g. RSA 2048/4096-bit, AES 256-bit) and active libraries."),
        ("Quantum-Safe Evaluation", "Classifies every primitive: flags RSA, ECC, and Diffie-Hellman as Quantum-Unsafe."),
        ("Post-Quantum Migration Path", "Advises engineering teams on transitioning to NIST FIPS post-quantum standards (ML-KEM / ML-DSA).")
    ], badge_color=COLOR_AMBER)

    add_card(s6, 6.8, 1.8, 5.6, 5.0, "Weak Cryptography Alerts with Fixes", [
        ("MD5 / SHA-1 Deprecation", "Detects obsolete collision-prone hashes with line number citation and SHA-256 upgrade code."),
        ("Disabled TLS Checks", "Flags verify=False and ssl.CERT_NONE that enable Man-in-the-Middle (MITM) attacks."),
        ("Predictable Randomness", "Detects standard random module used for security tokens and provides secrets.token_hex replacement."),
        ("Cryptographic Posture Score", "Rates total codebase cryptographic posture: Strong, Moderate, or Needs Attention.")
    ], badge_color=COLOR_ROSE)

    # =========================================================================
    # SLIDE 7: Pillar 3 — RCA with Recommended Code Fixes
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Pillar 3: Root Cause Analysis with Actionable Fixes", "DIAGNOSE -> EXPLAIN -> RECOMMEND FIX", "07/12")

    add_card(s7, 0.8, 1.8, 5.6, 5.0, "The 3-Step Remediation Paradigm", [
        ("1. Scan & Identify", "Pinpoints vulnerability with exact File Name and Line Number (e.g. ingest.py: Line 297)."),
        ("2. Explain Root Cause", "Explains WHY the vulnerability exists, the exploit mechanism, and contributing development factors."),
        ("3. Recommend Replacement Code", "Generates complete, drop-in replacement code snippet fixing the flaw."),
        ("Strict Safety Constraint", "The AI model NEVER modifies files directly; it recommends clean fixes for developer review.")
    ], badge_color=COLOR_ACCENT)

    add_card(s7, 6.8, 1.8, 5.6, 5.0, "Deep Architectural Checks Detected", [
        ("LLM Prompt Injection", "Flags unescaped context string interpolation and provides XML boundary tags (OWASP LLM01)."),
        ("Path Traversal Vulnerabilities", "Catches unvalidated user-supplied folder paths and supplies Path.resolve() guards."),
        ("Unauthenticated Databases", "Identifies unencrypted vector DB connections and recommends API-key/TLS client config."),
        ("Destructive Actions", "Flags unconditional collection drops and adds explicit confirmation flags."),
        ("Resource Exhaustion (DoS)", "Identifies unbounded read_text() calls and supplies file-size boundary checks.")
    ], badge_color=COLOR_GREEN)

    # =========================================================================
    # SLIDE 8: Pillar 4 — OWASP Top 10 & CWE Top 25
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Pillar 4: OWASP Top 10 & CWE Top 25 Gap Analysis", "STANDARDS COMPLIANCE", "08/12")

    add_card(s8, 0.8, 1.8, 5.6, 5.0, "OWASP Top 10 (2021) Assessment", [
        ("Full 10-Category Coverage", "Evaluates A01 Broken Access Control through A10 SSRF."),
        ("Visual Status Ratings", "Classifies each category: ✅ Covered, ⚠️ Partial, ❌ Gap, or ➖ Not Assessed."),
        ("Evidence & Residual Risk", "Documents specific line evidence and calculates residual risk (Critical, High, Medium, Low)."),
        ("Controls Observed", "Identifies existing defenses (parameterized queries, hashing, role guards).")
    ], badge_color=COLOR_ACCENT)

    add_card(s8, 6.8, 1.8, 5.6, 5.0, "CWE Top 25 & 4-Tier Roadmap", [
        ("CWE Vulnerability Mapping", "Correlates findings with CWE-798, CWE-78, CWE-89, CWE-95, CWE-502, CWE-22, CWE-306, CWE-327."),
        ("Priority 1 (Fix Immediately)", "Critical command injection, code evaluation, hardcoded credentials."),
        ("Priority 2 (Current Sprint)", "Disabled TLS verification, unsafe deserialization, broken access controls."),
        ("Priority 3 (Next Release)", "Predictable PRNGs, unpinned dependencies, quantum-unsafe ciphers."),
        ("Priority 4 (Ongoing / Hygiene)", "Sanitized error logging, security headers, CI/CD automated gates.")
    ], badge_color=COLOR_AMBER)

    # =========================================================================
    # SLIDE 9: CI/CD Pipeline & GitHub Composite Action
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Seamless DevSecOps: Automated CI/CD Integration", "AUTOMATION & GITHUB ACTIONS", "09/12")

    add_card(s9, 0.8, 1.8, 5.6, 5.0, "Ephemeral CI Runner Architecture", [
        ("Zero-Configuration Qdrant", "Spins up ephemeral Qdrant Docker container in seconds on GitHub Actions runner."),
        ("Direct Step Summary Publishing", "Appends reports directly to $GITHUB_STEP_SUMMARY — displays full tables without downloading zip files!"),
        ("Deterministic Fallback Guarantee", "If Qdrant or API keys are absent (e.g. fork PRs), deterministic mode runs with 100% success."),
        ("Quality Gate Enforcement", "Optional --fail-on-critical flag halts the build if CRITICAL security flaws are detected.")
    ], badge_color=COLOR_ACCENT)

    add_card(s9, 6.8, 1.8, 5.6, 5.0, "Reusable Composite Action (action.yml)", [
        ("One-Line Integration", "Any repository on GitHub can audit its codebase simply by adding:"),
        ("Action Syntax", "uses: menonkrishna57/Audit-App@main"),
        ("Custom Target Paths", "Supports auditing subdirectories, monorepos, or external checkouts via target-path input."),
        ("Optional AI Mode", "Pass google-api-key input to enable AI RAG; gracefully defaults to deterministic mode if omitted."),
        ("Automated GHCR Deployment", "Builds and pushes production multi-arch Docker images to GitHub Container Registry.")
    ], badge_color=COLOR_GREEN)

    # =========================================================================
    # SLIDE 10: Interactive Streamlit Web UI Dashboard
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Developer Experience: Interactive Security Dashboard", "STREAMLIT WEB UI", "10/12")

    add_card(s10, 0.8, 1.8, 3.6, 5.0, "Tab 1: Ingestion Engine", [
        ("Local Folder Picker", "Native Explorer dialog or text input to target any repository."),
        ("Multi-Provider Embeddings", "Toggle between Google Gemini, Jina AI Cloud, or Ollama local."),
        ("Real-time Status Monitor", "Visual indicators for Qdrant database, Google API key, and Jina key health.")
    ], badge_color=COLOR_ACCENT)

    add_card(s10, 4.8, 1.8, 3.6, 5.0, "Tab 2: Security Audits", [
        ("Quick Action One-Click", "Runs full pipeline: SBOM + CBOM + RCA + CWE Gap Analysis in one pass."),
        ("Granular Category Runs", "Inspect individual audit dimensions with custom top-k retrieval."),
        ("Gemini Model Selector", "Switch between gemini-3.5-flash, 3.5-flash-lite, 3.8-flash, or 3.5-pro on the fly.")
    ], badge_color=COLOR_GREEN)

    add_card(s10, 8.8, 1.8, 3.6, 5.0, "Tab 3: Past Reports", [
        ("Historical Audit Browser", "Instant browsing of all past generated audits stored in reports/."),
        ("Code Block Rendering", "Interactive syntax-highlighted code blocks with line numbers and recommended fixes."),
        ("Export & Download", "Download audit reports in clean Markdown (.md) or machine-readable JSON (.json).")
    ], badge_color=COLOR_AMBER)

    # =========================================================================
    # SLIDE 11: Competitive Advantage & Differentiation
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Why Audit-App Outperforms Traditional SAST Tools", "COMPETITIVE DIFFERENTIATION", "11/12")

    add_card(s11, 0.8, 1.8, 3.6, 5.0, "Traditional SAST (Sonar / Bandit)", [
        ("Rigid Pattern Matching", "High false positive rates on complex multi-chunk code logic."),
        ("No Cryptographic Posture", "Does not inventory quantum vulnerability or build a CBOM."),
        ("Generic Error Strings", "Tells developers what is wrong, but offers no verified code replacement."),
        ("Shallow Manifest Parsing", "Misses unpinned indirect imports and notebook cell dependencies.")
    ], badge_color=COLOR_ROSE)

    add_card(s11, 4.8, 1.8, 3.6, 5.0, "Generic AI Copilots", [
        ("Single-File Scope", "Copilots only view the current open buffer, missing repository-wide context."),
        ("No Compliance Structure", "Cannot produce formal SBOM, CBOM, or OWASP compliance matrices."),
        ("Hallucination Risk", "Lacks vector RAG verification grounding findings against real codebase chunks."),
        ("Not CI/CD Native", "Designed for autocomplete typing, not automated pipeline security gates.")
    ], badge_color=COLOR_AMBER)

    add_card(s11, 8.8, 1.8, 3.6, 5.0, "Audit-App Advantage", [
        ("Unified 4-Pillar Auditing", "Only platform combining SBOM, CBOM, RCA, and OWASP Gap Analysis natively."),
        ("Actionable Code Fixes", "Provides exact line numbers and copyable secure replacement code snippets."),
        ("Quantum-Safe Readiness", "Full inventory of crypto primitives with post-quantum migration advice."),
        ("Dual-Engine Resilience", "AI RAG for depth + deterministic fallback for fast, offline CI/CD.")
    ], badge_color=COLOR_GREEN)

    # =========================================================================
    # SLIDE 12: Summary & Future Roadmap
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Summary & Future Roadmap: Autonomous Code Governance", "CONCLUSION & VISION", "12/12")

    add_card(s12, 0.8, 1.8, 5.6, 5.0, "Current Achievements", [
        ("Production-Ready Core", "Complete 4-pillar security pipeline tested and verified with 100% pytest pass rate."),
        ("Language Versatility", "Deep analysis across Python (.py, .ipynb), Java (.java), C/C++ (.cpp, .c, .h), and manifests."),
        ("DevSecOps Native", "Turnkey GitHub Actions workflow + reusable composite action with $GITHUB_STEP_SUMMARY UI."),
        ("Dual AI & Deterministic Engine", "State-of-the-art Gemini LLM integration with zero-dependency offline continuity.")
    ], badge_color=COLOR_GREEN)

    add_card(s12, 6.8, 1.8, 5.6, 5.0, "Strategic Roadmap", [
        ("Autonomous PR Review Bot", "Auto-post line-level suggested code fixes directly onto GitHub Pull Request diffs."),
        ("SLSA Provenance Attestation", "Generate cryptographic SLSA build provenance attestations alongside SBOMs."),
        ("IDE Plugins", "Real-time audit diagnostics in VS Code and JetBrains IDEs as developers write code."),
        ("Enterprise Policy Engine", "Custom organizational compliance rulebooks and regulatory export (FedRAMP / ISO 27001).")
    ], badge_color=COLOR_ACCENT)

    out_file = Path(output_path).resolve()
    prs.save(str(out_file))
    print(f"[OK] Successfully generated 12-slide presentation at: {out_file}")
    return str(out_file)

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "Audit_App_Presentation.pptx"
    create_deck(out)
